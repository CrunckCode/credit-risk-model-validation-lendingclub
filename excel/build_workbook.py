"""Builds credit_validation_workbook.xlsx: live scorecard, deciles, calibration, PSI and expected-loss formulas over an OOT1 sample.

Typed numbers live on Inputs, Scorecard, Python_Ref, the typed Loan_Sample columns, and the typed PSI/EL reference columns.
Every other number is a formula. Run via: py -3 -m credit_validation.reconcile_excel (build + LibreOffice recalculation).
"""
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "python"))
from credit_validation import config, stability, validation  # noqa: E402
from credit_validation import columns as C  # noqa: E402
from credit_validation import reconcile_excel as rx  # noqa: E402

# ===== CONFIG (user inputs) =====
OUT_FILE = Path(__file__).resolve().parent / "credit_validation_workbook.xlsx"
AUTHOR = "Deepak Chaudhary"
HDR, FIRST = 4, 5                      # Loan_Sample header row and first data row
CALC_BASE = {C.SUB_GRADE: "B2", C.HOME_OWNERSHIP: "RENT", C.ANNUAL_INC: 85000.0, C.PURPOSE: "credit_card",
             C.ADDR_STATE: "FL", C.DTI: 19.29, C.INQ_LAST_6MTHS: 1.0, C.REVOL_UTIL: 57.4, C.LOAN_TO_INC: 0.0823529}
CALC_SHOCK_VAR, CALC_SHOCK_PTS = C.DTI, rx.CALC_DTI_SHOCK
HL_DF_OOT = V_N = validation.N_BINS    # OOT sample: df = g, as in validation.hosmer_lemeshow(in_sample=False)
RESULT_COLS = ("test", "split", "value", "light")
FAR = 1e300                            # stands in for +/- infinity in the PSI bin helpers
POP_SPLIT = C.OOT1
# ===== END CONFIG =====

BOLD = Font(bold=True)
_CACHE = {}


def calc_cases() -> dict:
    shocked = dict(CALC_BASE)
    shocked[CALC_SHOCK_VAR] = CALC_BASE[CALC_SHOCK_VAR] + CALC_SHOCK_PTS
    return {"base": dict(CALC_BASE), f"{CALC_SHOCK_VAR}_plus_{CALC_SHOCK_PTS:g}": shocked}


def _ctx(quick):
    if quick not in _CACHE:
        sample, pts, psi_bins, realized = rx.load_tables(quick)
        sc = rx.load_scorecard()
        ref = rx.sample_reference(sample, sc, psi_bins)
        res = pd.read_csv(rx.RESULTS_PATH)
        res = res[res["model"] == rx.MODEL_NAME].reset_index(drop=True)
        _CACHE[quick] = (sample, pts, psi_bins, realized, sc, ref, res)
    return _CACHE[quick]


def _val(x):
    """openpyxl-safe python scalar (None for NaN)."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return None
    return x.item() if hasattr(x, "item") else x


def build(out_file=OUT_FILE, calc_inputs=None, quick=False):
    sample, pts, psi_bins, realized, sc, ref, res = _ctx(quick)
    calc_inputs = calc_inputs or CALC_BASE
    variables = list(sc.variables)
    n = len(sample)
    last = FIRST + n - 1
    wb = openpyxl.Workbook()
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    names = ["README", "Inputs", "Scorecard", "Score_Calc", "Loan_Sample", "Deciles", "Calibration", "PSI", "EL",
             "Checks", "Conclusions", "Python_Ref"]
    ws = {nm: (wb.active if i == 0 else wb.create_sheet(nm)) for i, nm in enumerate(names)}
    ws["README"].title = "README"
    addr = {}

    def reg(name, sheet, ref_):
        addr[name] = f"{sheet}!{ref_}"

    def head(sh, row, labels, col0=1):
        for j, t in enumerate(labels):
            sh.cell(row, col0 + j, t).font = BOLD

    # ============================================================ Inputs
    wi = ws["Inputs"]
    NM = {}
    row = [3]

    def put(key, label, value, note=""):
        r = row[0]
        wi.cell(r, 1, label)
        wi.cell(r, 2, value)
        wi.cell(r, 3, note)
        NM[key] = f"Inputs!$B${r}"
        row[0] += 1

    wi["A1"] = "Inputs (all typed model parameters, thresholds, tolerances and calculator attributes live here)"
    wi["A1"].font = BOLD
    put("pdo", "Scorecard PDO (points to double the odds)", config.PDO, "config.PDO")
    put("base_score", "Scorecard base score", config.BASE_SCORE, "config.BASE_SCORE")
    put("base_odds", "Scorecard base odds (good:bad)", config.BASE_ODDS, "config.BASE_ODDS")
    put("factor", "Scorecard factor", float(sc.factor), "Typed from the fitted scorecard so ties and bin edges match Python bit for bit")
    put("offset", "Scorecard offset", float(sc.offset), "Typed from the fitted scorecard")
    put("factor_chk", "Factor recomputed from PDO", f"={NM['pdo']}/LN(2)", "PDO / ln 2 (formula, compared on Checks)")
    put("offset_chk", "Offset recomputed from base score and odds", f"={NM['base_score']}-{NM['factor_chk']}*LN({NM['base_odds']})",
        "base score - factor x ln(base odds) (formula, compared on Checks)")
    put("intercept", "Scorecard logit intercept", float(sc.intercept), "Fitted scorecard")
    put("n_dec", "Number of deciles / HL groups", validation.N_BINS, "validation.N_BINS")
    put("hl_df", "Hosmer-Lemeshow degrees of freedom", HL_DF_OOT, "df = g on out-of-time data (validation.hosmer_lemeshow)")
    put("psi_floor", "PSI share floor", stability.PSI_FLOOR, "stability.PSI_FLOOR")
    put("psi_green", "PSI green limit", config.THRESHOLDS["psi"][0], "config.THRESHOLDS")
    put("psi_amber", "PSI amber limit", config.THRESHOLDS["psi"][1], "config.THRESHOLDS")
    for key in ("gini", "ks", "hl_p", "binom_p", "central_tendency"):
        put(f"{key}_green", f"{key} green limit", config.THRESHOLDS[key][0], "config.THRESHOLDS")
        put(f"{key}_amber", f"{key} amber limit", config.THRESHOLDS[key][1], "config.THRESHOLDS")
    put("tol_pd", "Tolerance: PD and logit", rx.TOL_PD, "Checks sheet")
    put("tol_pts", "Tolerance: points and score", rx.TOL_POINTS, "Checks sheet")
    put("tol_stat", "Tolerance: HL statistic", rx.TOL_HL_STAT, "Checks sheet")
    put("tol_p", "Tolerance: p-values", rx.TOL_PVALUE, "Checks sheet")
    put("tol_rate", "Tolerance: rates, KS, PSI", rx.TOL_RATE, "Checks sheet")
    put("tol_count", "Tolerance: counts", 0, "Checks sheet (exact)")
    put("tol_usd", "Tolerance: expected loss (USD)", rx.TOL_EL_USD, "Checks sheet")
    row[0] += 1
    head(wi, row[0], ["One-loan calculator attributes (clear a cell to test the missing bin)", "value"])
    row[0] += 1
    for v in variables:
        wi.cell(row[0], 1, v)
        wi.cell(row[0], 2, calc_inputs.get(v))
        NM[f"calc_{v}"] = f"Inputs!$B${row[0]}"
        row[0] += 1
    wi.column_dimensions["A"].width = 58
    wi.column_dimensions["C"].width = 80

    # ============================================================ Scorecard
    wsc = ws["Scorecard"]
    wsc["A1"] = f"Scorecard {rx.MODEL_NAME}: typed points table (scorecard_points.csv). Bins are (lo, hi]; missing is the last bin of each variable."
    wsc["A1"].font = BOLD
    wsc["A2"] = "Column I holds the finite upper edges used by the bin lookup. Columns K to P summarise each variable; R to T map each category to its bin."
    head(wsc, 4, [rx.VAR_COL, rx.BIN_COL, rx.LO_COL, rx.HI_COL, rx.CAT_COL, rx.WOE_COL, rx.COEF_COL, rx.POINTS_COL, "edge_helper"])
    head(wsc, 4, ["variable", "kind", "n_bins", "coef", "default_bin", "first_row"], col0=11)
    head(wsc, 4, ["cat_variable", "category", "bin_number"], col0=18)
    SC = {}
    r = 5
    cat_r = 5
    for vi, v in enumerate(variables):
        g = pts[pts[rx.VAR_COL] == v].reset_index(drop=True)
        spec = sc.specs[v]
        nb = len(g)
        assert nb == spec.n_bins, f"{v}: points table and scorecard disagree"
        is_cat = spec.kind == "categorical"
        r0 = r
        for i, rec in g.iterrows():
            wsc.cell(r, 1, v)
            wsc.cell(r, 2, int(rec[rx.BIN_COL]))
            lo, hi = rec[rx.LO_COL], rec[rx.HI_COL]
            if not pd.isna(lo):
                wsc.cell(r, 3, "-inf" if np.isinf(lo) else float(lo))
                wsc.cell(r, 4, "inf" if np.isinf(hi) else float(hi))
                if not np.isinf(hi):
                    wsc.cell(r, 9, float(hi))
            if isinstance(rec[rx.CAT_COL], str):
                wsc.cell(r, 5, rec[rx.CAT_COL])
                for c in rec[rx.CAT_COL].split(rx.CAT_JOIN):
                    wsc.cell(cat_r, 18, v)
                    wsc.cell(cat_r, 19, c)
                    wsc.cell(cat_r, 20, int(rec[rx.BIN_COL]) + 1)
                    cat_r += 1
            wsc.cell(r, 6, float(rec[rx.WOE_COL]))
            wsc.cell(r, 7, float(rec[rx.COEF_COL]))
            wsc.cell(r, 8, float(rec[rx.POINTS_COL]))
            r += 1
        d = dict(nb=nb, is_cat=is_cat, woe=f"Scorecard!$F${r0}:$F${r0 + nb - 1}", pts=f"Scorecard!$H${r0}:$H${r0 + nb - 1}",
                 coef=f"Scorecard!$N${5 + vi}", default=f"Scorecard!$O${5 + vi}")
        if is_cat:
            cr0 = cat_r - sum(len(str(x).split(rx.CAT_JOIN)) for x in g[rx.CAT_COL] if isinstance(x, str))
            d["cats"] = f"Scorecard!$S${cr0}:$S${cat_r - 1}"
            d["catbin"] = f"Scorecard!$T${cr0}:$T${cat_r - 1}"
        else:
            assert nb >= 3, f"{v}: lookup needs at least 2 regular bins"
            d["edges"] = f"Scorecard!$I${r0}:$I${r0 + nb - 3}"
        SC[v] = d
        rr = 5 + vi
        wsc.cell(rr, 11, v)
        wsc.cell(rr, 12, spec.kind)
        wsc.cell(rr, 13, nb)
        wsc.cell(rr, 14, float(sc.coefs[v]))
        wsc.cell(rr, 15, int(spec.default_group) + 1 if is_cat else None)
        wsc.cell(rr, 16, r0)
    sr = 5 + len(variables) + 2
    wsc.cell(sr, 11, "intercept")
    wsc.cell(sr, 12, f"={NM['intercept']}")
    wsc.cell(sr + 1, 11, "factor")
    wsc.cell(sr + 1, 12, f"={NM['factor']}")
    wsc.cell(sr + 2, 11, "offset")
    wsc.cell(sr + 2, 12, f"={NM['offset']}")
    wsc.column_dimensions["A"].width = 18
    wsc.column_dimensions["E"].width = 40

    def bin_expr(v, x):
        d = SC[v]
        if d["is_cat"]:
            return f"IF(ISBLANK({x}),{d['nb']},IFERROR(INDEX({d['catbin']},MATCH({x},{d['cats']},0)),{d['default']}))"
        return f"IF(ISBLANK({x}),{d['nb']},SUMPRODUCT(--({d['edges']}<{x}))+1)"

    # ============================================================ Loan_Sample
    wl = ws["Loan_Sample"]
    wl["A1"] = f"Loan sample: {n:,} out-of-time (2013) loans stratified by grade. Typed columns come from Python; xl_ columns are live formulas."
    wl["A1"].font = BOLD
    wl["A2"] = "Attributes are stored at full precision (float32 values widened exactly) so values on a bin edge land in the same bin as Python."
    static = list(sample.columns)
    live = ([f"xl_bin_{v}" for v in variables] + [f"xl_woe_{v}" for v in variables] + [f"xl_pts_{v}" for v in variables] +
            ["xl_margin", "xl_score", "xl_score_sumpts", "xl_pd", "xl_el", "xl_rank_desc", "xl_decile", "xl_rank_asc", "xl_cal_bin",
             "xl_pd_diff", "xl_score_diff"])
    lc = {}
    for j, h in enumerate(static + live, start=1):
        wl.cell(HDR, j, h).font = BOLD
        lc[h] = L(j)
    vals = {h: sample[h].to_numpy(dtype=object) for h in static}
    for i in range(n):
        rr = FIRST + i
        for h in static:
            v = _val(vals[h][i])
            if v is not None:
                wl[f"{lc[h]}{rr}"] = v

    def rng(h):
        return f"${lc[h]}${FIRST}:${lc[h]}${last}"

    def lsr(h):
        return f"Loan_Sample!{rng(h)}"

    pdl, ranka = lc["xl_pd"], lc["xl_rank_asc"]
    for i in range(n):
        rr = FIRST + i
        for v in variables:
            wl[f"{lc['xl_bin_' + v]}{rr}"] = "=" + bin_expr(v, f"{lc[v]}{rr}")
            b = f"{lc['xl_bin_' + v]}{rr}"
            wl[f"{lc['xl_woe_' + v]}{rr}"] = f"=INDEX({SC[v]['woe']},{b})"
            wl[f"{lc['xl_pts_' + v]}{rr}"] = f"=INDEX({SC[v]['pts']},{b})"
        m = f"{lc['xl_margin']}{rr}"
        wl[m] = "=" + NM["intercept"] + "".join(f"+{SC[v]['coef']}*{lc['xl_woe_' + v]}{rr}" for v in variables)
        wl[f"{lc['xl_score']}{rr}"] = f"={NM['offset']}-{NM['factor']}*{m}"
        wl[f"{lc['xl_score_sumpts']}{rr}"] = f"=SUM({lc['xl_pts_' + variables[0]]}{rr}:{lc['xl_pts_' + variables[-1]]}{rr})"
        wl[f"{pdl}{rr}"] = f"=1/(1+EXP(-{m}))"
        wl[f"{lc['xl_el']}{rr}"] = f"={pdl}{rr}*{lc[C.PRED_LGD]}{rr}*{lc[C.PRED_EAD]}{rr}"
        prng = rng("xl_pd")
        for out, how in (("xl_rank_desc", 0), ("xl_rank_asc", 1)):
            wl[f"{lc[out]}{rr}"] = f"=RANK({pdl}{rr},{prng},{how})+COUNTIF(${pdl}${HDR}:{pdl}{rr - 1},{pdl}{rr})"
        wl[f"{lc['xl_decile']}{rr}"] = f"=INT(({lc['xl_rank_desc']}{rr}-1)*{NM['n_dec']}/Deciles!$B$3)+1"
        wl[f"{lc['xl_cal_bin']}{rr}"] = f"=INT(({ranka}{rr}-1)*{NM['n_dec']}/Deciles!$B$3)+1"
        wl[f"{lc['xl_pd_diff']}{rr}"] = f"=ABS({pdl}{rr}-{lc[C.PRED_PD]}{rr})"
        wl[f"{lc['xl_score_diff']}{rr}"] = f"=ABS({lc['xl_score']}{rr}-{lc[C.SCORE]}{rr})"
    wl.freeze_panes = f"B{FIRST}"

    # ============================================================ Score_Calc
    wc = ws["Score_Calc"]
    wc["A1"] = "One-loan scorecard calculator (attributes are on Inputs; a cleared input falls in the missing bin)"
    wc["A1"].font = BOLD
    CF = 6
    head(wc, CF - 1, ["variable", "value", "bin_number", "woe", "coef", "coef x woe", "points"])
    c1 = CF + len(variables) - 1
    for i, v in enumerate(variables):
        r_ = CF + i
        d = SC[v]
        wc.cell(r_, 1, v)
        wc.cell(r_, 2, f'=IF(ISBLANK({NM["calc_" + v]}),"",{NM["calc_" + v]})')
        if d["is_cat"]:
            wc.cell(r_, 3, f'=IF(B{r_}="",{d["nb"]},IFERROR(INDEX({d["catbin"]},MATCH(B{r_},{d["cats"]},0)),{d["default"]}))')
        else:
            wc.cell(r_, 3, f'=IF(B{r_}="",{d["nb"]},SUMPRODUCT(--({d["edges"]}<B{r_}))+1)')
        wc.cell(r_, 4, f"=INDEX({d['woe']},C{r_})")
        wc.cell(r_, 5, f"={d['coef']}")
        wc.cell(r_, 6, f"=D{r_}*E{r_}")
        wc.cell(r_, 7, f"=INDEX({d['pts']},C{r_})")
    rs = c1 + 2
    summary = [("Sum of coef x woe", f"=SUM(F{CF}:F{c1})", None), ("Logit (margin)", f"={NM['intercept']}+B{rs}", "calc_margin"),
               ("PD = 1/(1+EXP(-logit))", f"=1/(1+EXP(-B{rs + 1}))", "calc_pd"),
               ("Total score = offset - factor x logit", f"={NM['offset']}-{NM['factor']}*B{rs + 1}", "calc_score"),
               ("Sum of per-variable points (equals total score)", f"=SUM(G{CF}:G{c1})", "calc_points_sum")]
    for off, (lab, f, key) in enumerate(summary):
        wc.cell(rs + off, 1, lab)
        wc.cell(rs + off, 2, f)
        if key:
            reg(key, "Score_Calc", f"B{rs + off}")
    sg = rs + 7
    sgv = C.SUB_GRADE
    sg_row = CF + variables.index(sgv) if sgv in variables else None
    wc.cell(sg - 1, 1, "Sub-grade comparison (sample)").font = BOLD
    if sg_row is not None:
        gcol, ycol = lsr(sgv), lsr(C.TARGET_BAD12)
        items = [("Sub-grade entered", f"=B{sg_row}"),
                 ("Loans in sample with this sub-grade", f"=COUNTIF({gcol},B{sg})"),
                 ("Observed bad rate of this sub-grade in sample", f'=IFERROR(SUMIFS({ycol},{gcol},B{sg})/B{sg + 1},"")'),
                 ("Mean scorecard PD of this sub-grade in sample", f'=IFERROR(SUMIFS({lsr("xl_pd")},{gcol},B{sg})/B{sg + 1},"")'),
                 ("Calculator PD", f"=B{rs + 2}"),
                 ("Calculator PD minus sub-grade mean PD", f'=IFERROR(B{sg + 4}-B{sg + 3},"")'),
                 ("Points from the sub-grade bin alone", f"=G{sg_row}")]
        for k, (lab, f) in enumerate(items):
            wc.cell(sg + k, 1, lab)
            wc.cell(sg + k, 2, f)
    wc.column_dimensions["A"].width = 52
    reg("calc_first", "Score_Calc", f"A{CF}")

    # ============================================================ Deciles
    wd = ws["Deciles"]
    wd["A1"] = "Decile table over the live PD: rank descending by PD, ties broken by input order, decile = INT((rank-1)*10/N)+1"
    wd["A1"].font = BOLD
    wd["A3"] = "Loans (N)"
    wd["B3"] = f"=COUNT({lsr('xl_pd')})"
    wd["A4"] = "Bads"
    wd["B4"] = f"=SUM({lsr(C.TARGET_BAD12)})"
    wd["A5"] = "Overall bad rate"
    wd["B5"] = "=B4/B3"
    DH = 8
    D0, D1 = DH + 1, DH + validation.N_BINS
    head(wd, DH, ["decile", "n", "bads", "bad_rate", "mean_pd", "cum_bads", "cum_bad_share (capture)", "cum_good_share", "ks_at_decile",
                  "lift", "cum_lift", "roc_trapezoid_area"])
    dec, ybad = lsr("xl_decile"), lsr(C.TARGET_BAD12)
    for i in range(validation.N_BINS):
        r_ = D0 + i
        wd.cell(r_, 1, f"=ROW()-{DH}")
        wd.cell(r_, 2, f"=COUNTIFS({dec},A{r_})")
        wd.cell(r_, 3, f"=SUMIFS({ybad},{dec},A{r_})")
        wd.cell(r_, 4, f"=C{r_}/B{r_}")
        wd.cell(r_, 5, f"=SUMIFS({lsr('xl_pd')},{dec},A{r_})/B{r_}")
        wd.cell(r_, 6, f"=SUM(C${D0}:C{r_})")
        wd.cell(r_, 7, f"=F{r_}/$B$4")
        wd.cell(r_, 8, f"=(SUM(B${D0}:B{r_})-F{r_})/($B$3-$B$4)")
        wd.cell(r_, 9, f"=ABS(G{r_}-H{r_})")
        wd.cell(r_, 10, f"=D{r_}/$B$5")
        wd.cell(r_, 11, f"=(F{r_}/SUM(B${D0}:B{r_}))/$B$5")
        pg, ph = (f"G{r_ - 1}", f"H{r_ - 1}") if i else ("0", "0")
        wd.cell(r_, 12, f"=(H{r_}-{ph})*(G{r_}+{pg})/2")
    wd.cell(D1 + 2, 1, "Max KS (decile based)")
    wd.cell(D1 + 2, 2, f"=MAX(I{D0}:I{D1})")
    reg("ks_max", "Deciles", f"B{D1 + 2}")
    wd.cell(D1 + 3, 1, "AUC, approximate (ROC trapezoid over deciles)")
    wd.cell(D1 + 3, 2, f"=SUM(L{D0}:L{D1})")
    wd.cell(D1 + 4, 1, "Gini, approximate (2 x AUC - 1; exact Gini is on Python_Ref)")
    wd.cell(D1 + 4, 2, f"=2*B{D1 + 3}-1")
    reg("gini_approx", "Deciles", f"B{D1 + 4}")
    wd.column_dimensions["A"].width = 52

    # ============================================================ Calibration
    wk = ws["Calibration"]
    wk["A1"] = "Calibration: Hosmer-Lemeshow groups (ascending PD, ties by input order), grade tests, central tendency, Brier"
    wk["A1"].font = BOLD
    KH = 5
    K0, K1 = KH + 1, KH + validation.N_BINS
    head(wk, KH, ["hl_group", "n", "mean_pd", "expected_bads", "observed_bads", "observed_rate", "hl_term", "obs_minus_mean_pd"])
    cb = lsr("xl_cal_bin")
    for i in range(validation.N_BINS):
        r_ = K0 + i
        wk.cell(r_, 1, f"=ROW()-{KH}")
        wk.cell(r_, 2, f"=COUNTIFS({cb},A{r_})")
        wk.cell(r_, 3, f"=SUMIFS({lsr('xl_pd')},{cb},A{r_})/B{r_}")
        wk.cell(r_, 4, f"=B{r_}*C{r_}")
        wk.cell(r_, 5, f"=SUMIFS({ybad},{cb},A{r_})")
        wk.cell(r_, 6, f"=E{r_}/B{r_}")
        wk.cell(r_, 7, f"=(E{r_}-D{r_})^2/(B{r_}*C{r_}*(1-C{r_}))")
        wk.cell(r_, 8, f"=F{r_}-C{r_}")
    s_ = K1 + 2
    cal_items = [("hl_stat", "Hosmer-Lemeshow statistic", f"=SUM(G{K0}:G{K1})"), ("hl_df", "HL degrees of freedom", f"={NM['hl_df']}"),
                 ("hl_p", "HL p-value (right tail of chi-square)", f"=CHIDIST(B{s_},B{s_ + 1})"),
                 ("brier", "Brier score", f"=SUMPRODUCT(({lsr('xl_pd')}-{ybad})^2)/Deciles!$B$3"),
                 ("brier_skill", "Brier skill vs base rate", f"=1-B{s_ + 3}/(Deciles!$B$5*(1-Deciles!$B$5))"),
                 ("mean_pd", "Mean predicted PD", f"=AVERAGE({lsr('xl_pd')})"),
                 ("obs_rate", "Observed bad rate", "=Deciles!$B$5"),
                 ("ct_ratio", "Central tendency ratio (mean PD / observed)", f"=B{s_ + 5}/B{s_ + 6}"),
                 ("ct_rel_dev", "Central tendency relative deviation |mean PD - observed| / observed", f"=ABS(B{s_ + 5}-B{s_ + 6})/B{s_ + 6}")]
    for k, (key, lab, f) in enumerate(cal_items):
        wk.cell(s_ + k, 1, lab)
        wk.cell(s_ + k, 2, f)
        reg(key, "Calibration", f"B{s_ + k}")
    sl = s_ + len(cal_items) + 1
    wk.cell(sl, 1, "Calibration slope, sample (Python, static)")
    GH = sl + 4
    G0 = GH + 1
    grades = sorted(sample[C.GRADE].astype(str).unique())
    head(wk, GH, ["grade", "n", "bads", "observed_rate", "mean_pd", "binomial_z (PD understated)", "p_value_z (one-sided)",
                  "p_value_jeffreys (one-sided)"])
    gcol = lsr(C.GRADE)
    for i, g in enumerate(grades):
        r_ = G0 + i
        wk.cell(r_, 1, g)
        wk.cell(r_, 2, f"=COUNTIF({gcol},A{r_})")
        wk.cell(r_, 3, f"=SUMIFS({ybad},{gcol},A{r_})")
        wk.cell(r_, 4, f"=C{r_}/B{r_}")
        wk.cell(r_, 5, f"=SUMIFS({lsr('xl_pd')},{gcol},A{r_})/B{r_}")
        wk.cell(r_, 6, f"=(C{r_}-B{r_}*E{r_})/SQRT(B{r_}*E{r_}*(1-E{r_}))")
        wk.cell(r_, 7, f"=NORMSDIST(-F{r_})")
        wk.cell(r_, 8, f"=BETADIST(E{r_},C{r_}+0.5,B{r_}-C{r_}+0.5)")
    wk.column_dimensions["A"].width = 60
    cal_sl_cell = (wk, sl)

    # ============================================================ PSI
    wp = ws["PSI"]
    wp["A1"] = "Score PSI: typed development bins (dev_psi_bins.csv, columns A to G) against live out-of-time shares (columns H to L)"
    wp["A1"].font = BOLD
    wp["A2"] = "Bins are (lo, hi]. Columns F and G hold the same edges with +/-1E+300 for infinity so the live counts compare numerically."
    PH = 5
    P0 = PH + 1
    P1 = PH + len(psi_bins)
    head(wp, PH, ["bin", "lo", "hi", "dev_share", "dev_count", "lo_calc", "hi_calc", "oot_count", "oot_share", "dev_share_floored",
                  "oot_share_floored", "psi_contribution"])
    score_r = lsr("xl_score")
    for i, rec in psi_bins.reset_index(drop=True).iterrows():
        r_ = P0 + i
        wp.cell(r_, 1, i + 1)
        lo, hi = float(rec[rx.PSI_LO_COL]), float(rec[rx.PSI_HI_COL])
        wp.cell(r_, 2, "-inf" if np.isinf(lo) else lo)
        wp.cell(r_, 3, "inf" if np.isinf(hi) else hi)
        wp.cell(r_, 4, float(rec[rx.PSI_SHARE_COL]))
        wp.cell(r_, 5, int(rec[rx.PSI_COUNT_COL]))
        wp.cell(r_, 6, -FAR if np.isinf(lo) else lo)
        wp.cell(r_, 7, FAR if np.isinf(hi) else hi)
        wp.cell(r_, 8, f"=SUMPRODUCT(({score_r}>F{r_})*({score_r}<=G{r_}))")
        wp.cell(r_, 9, f"=H{r_}/Deciles!$B$3")
        wp.cell(r_, 10, f"=MAX(D{r_},{NM['psi_floor']})")
        wp.cell(r_, 11, f"=MAX(I{r_},{NM['psi_floor']})")
        wp.cell(r_, 12, f"=(K{r_}-J{r_})*LN(K{r_}/J{r_})")
    wp.cell(P1 + 2, 1, "Total PSI")
    wp.cell(P1 + 2, 2, f"=SUM(L{P0}:L{P1})")
    reg("psi_total", "PSI", f"B{P1 + 2}")
    wp.cell(P1 + 3, 1, "Loans binned (equals N)")
    wp.cell(P1 + 3, 2, f"=SUM(H{P0}:H{P1})")
    wp.cell(P1 + 4, 1, "At or below green limit (1 = yes)")
    wp.cell(P1 + 4, 2, f"=--(B{P1 + 2}<={NM['psi_green']})")
    wp.cell(P1 + 5, 1, "Above green and at or below amber limit (1 = yes)")
    wp.cell(P1 + 5, 2, f"=--AND(B{P1 + 2}>{NM['psi_green']},B{P1 + 2}<={NM['psi_amber']})")
    wp.column_dimensions["A"].width = 48

    # ============================================================ EL
    we = ws["EL"]
    we["A1"] = "Expected loss: live sample EL = PD x LGD x EAD by grade against realized 12-month net loss (typed columns H to J are OOT1 population values)"
    we["A1"].font = BOLD
    we["A2"] = "Scaled EL multiplies the sample EL by population funded / sample funded. Latest status date is Jan 2016, so 2013 outcomes are essentially resolved."
    EH = 5
    E0 = EH + 1
    head(we, EH, ["grade", "n", "funded", "mean_pd", "el_sample", "realized_loss_sample", "el_to_realized_sample", "oot1_n_typed",
                  "oot1_funded_typed", "oot1_realized_typed", "el_scaled_to_population", "el_scaled_to_realized_population",
                  "el_pct_funded", "realized_pct_funded"])
    rl = realized.set_index(C.GRADE)
    for i, g in enumerate(grades):
        r_ = E0 + i
        we.cell(r_, 1, g)
        we.cell(r_, 2, f"=COUNTIF({gcol},A{r_})")
        we.cell(r_, 3, f"=SUMIFS({lsr(C.FUNDED_AMNT)},{gcol},A{r_})")
        we.cell(r_, 4, f"=SUMIFS({lsr('xl_pd')},{gcol},A{r_})/B{r_}")
        we.cell(r_, 5, f"=SUMIFS({lsr('xl_el')},{gcol},A{r_})")
        we.cell(r_, 6, f"=SUMIFS({lsr(rx.REALIZED_COL)},{gcol},A{r_})")
        we.cell(r_, 7, f'=IFERROR(E{r_}/F{r_},"")')
        we.cell(r_, 8, float(rl.loc[g, "oot1_n"]))
        we.cell(r_, 9, float(rl.loc[g, "oot1_funded_amnt"]))
        we.cell(r_, 10, float(rl.loc[g, "oot1_realized"]))
        we.cell(r_, 11, f"=E{r_}*I{r_}/C{r_}")
        we.cell(r_, 12, f"=K{r_}/J{r_}")
        we.cell(r_, 13, f"=E{r_}/C{r_}")
        we.cell(r_, 14, f"=F{r_}/C{r_}")
    t_ = E0 + len(grades)
    we.cell(t_, 1, "Total")
    for col in "BCEFHIJK":
        we[f"{col}{t_}"] = f"=SUM({col}{E0}:{col}{t_ - 1})"
    we[f"D{t_}"] = f"=AVERAGE({lsr('xl_pd')})"
    we[f"G{t_}"] = f'=IFERROR(E{t_}/F{t_},"")'
    we[f"L{t_}"] = f"=K{t_}/J{t_}"
    we[f"M{t_}"] = f"=E{t_}/C{t_}"
    we[f"N{t_}"] = f"=F{t_}/C{t_}"
    reg("el_total", "EL", f"E{t_}")
    reg("el_realized_total", "EL", f"F{t_}")
    reg("el_ratio_pop", "EL", f"L{t_}")
    we.column_dimensions["A"].width = 14

    # ============================================================ Python_Ref
    wr = ws["Python_Ref"]
    wr["A1"] = "Python reference values (static). Block A: sample recomputation with the project modules. Block B: full OOT1 population from validation_results.csv."
    wr["A1"].font = BOLD
    head(wr, 3, ["metric", "value", "note"])
    cen, hl = ref["central"], ref["hl"]
    ref_items = [("ref_n", "Sample loans", ref["n"]), ("ref_bads", "Sample bads", ref["bads"]), ("ref_bad_rate", "Sample bad rate", ref["bad_rate"]),
                 ("ref_mean_pd", "Sample mean PD", ref["mean_pd"]), ("ref_brier", "Sample Brier", ref["brier"]),
                 ("ref_ks_decile", "Sample KS (decile based)", ref["ks_decile"]), ("ref_gini_approx", "Sample Gini approx (decile trapezoid)", ref["gini_approx"]),
                 ("ref_hl_stat", "Sample Hosmer-Lemeshow statistic", hl["stat"]), ("ref_hl_p", "Sample HL p-value", hl["p_value"]),
                 ("ref_ct_ratio", "Sample central tendency ratio", cen["ratio"]), ("ref_psi_total", "Sample score PSI", ref["psi_total"]),
                 ("ref_el_total", "Sample total expected loss (USD)", float(ref["el"].sum())),
                 ("ref_auc", "Sample AUC (exact)", ref["auc"]), ("ref_gini", "Sample Gini (exact)", ref["gini"]), ("ref_ks", "Sample KS (exact, loan level)", ref["ks"]),
                 ("ref_slope", "Sample calibration slope", ref["slope"]["slope"]), ("ref_intercept", "Sample calibration intercept", ref["slope"]["intercept"]),
                 ("ref_zero", "Zero (exported-column differences should be zero)", 0.0)]
    for k, (key, lab, v) in enumerate(ref_items):
        wr.cell(4 + k, 1, lab)
        wr.cell(4 + k, 2, float(v))
        reg(key, "Python_Ref", f"B{4 + k}")
    pr = 4 + len(ref_items) + 2
    wr.cell(pr, 1, f"Block B: {rx.MODEL_NAME} validation_results.csv (full population, not the sample)").font = BOLD
    head(wr, pr + 1, list(RESULT_COLS))
    for k, rec in res.iterrows():
        for j, c in enumerate(RESULT_COLS, start=1):
            v = _val(rec[c])
            if v is not None:
                wr.cell(pr + 2 + k, j, v)
    pop = res[res["split"] == POP_SPLIT].set_index("test")["value"]
    wr.column_dimensions["A"].width = 52
    slope_row = 4 + [k for k, it in enumerate(ref_items) if it[0] == "ref_slope"][0]
    cal_sl_cell[0].cell(cal_sl_cell[1], 2, f"=Python_Ref!$B${slope_row}")
    cal_sl_cell[0].cell(cal_sl_cell[1] + 1, 1, "Calibration slope, full OOT1 population (Python, static)")
    pop_slope_row = pr + 2 + int(res.index[(res["test"] == "cal_slope") & (res["split"] == POP_SPLIT)][0])
    cal_sl_cell[0].cell(cal_sl_cell[1] + 1, 2, f"=Python_Ref!$C${pop_slope_row}")

    # ============================================================ Checks
    wx = ws["Checks"]
    wx["A1"] = "Checks: live workbook values against Python_Ref (absolute differences)"
    wx["A1"].font = BOLD
    wx["A2"] = "All checks pass (1 = yes)"
    head(wx, 4, ["check", "live", "python", "abs_diff", "tolerance", "pass"])
    def R(k):
        return "=" + addr[k]

    checks = [("Sample loans", f"=Deciles!$B$3", R("ref_n"), "tol_count"),
              ("Sample bads", "=Deciles!$B$4", R("ref_bads"), "tol_count"),
              ("Bad rate", "=Deciles!$B$5", R("ref_bad_rate"), "tol_rate"),
              ("Mean PD", "=" + addr["mean_pd"], R("ref_mean_pd"), "tol_pd"),
              ("Brier score", "=" + addr["brier"], R("ref_brier"), "tol_rate"),
              ("KS (decile based)", "=" + addr["ks_max"], R("ref_ks_decile"), "tol_rate"),
              ("Gini approx (decile trapezoid)", "=" + addr["gini_approx"], R("ref_gini_approx"), "tol_rate"),
              ("Hosmer-Lemeshow statistic", "=" + addr["hl_stat"], R("ref_hl_stat"), "tol_stat"),
              ("Hosmer-Lemeshow p-value", "=" + addr["hl_p"], R("ref_hl_p"), "tol_p"),
              ("Central tendency ratio", "=" + addr["ct_ratio"], R("ref_ct_ratio"), "tol_rate"),
              ("Score PSI", "=" + addr["psi_total"], R("ref_psi_total"), "tol_rate"),
              ("Loans binned in PSI", f"=PSI!$B${P1 + 3}", R("ref_n"), "tol_count"),
              ("Total expected loss (USD)", "=" + addr["el_total"], R("ref_el_total"), "tol_usd"),
              ("Max PD difference vs exported Python PD", f"=MAX({lsr('xl_pd_diff')})", R("ref_zero"), "tol_pd"),
              ("Max score difference vs exported Python score", f"=MAX({lsr('xl_score_diff')})", R("ref_zero"), "tol_pts"),
              ("Sum of abs differences, sum of points vs offset - factor x logit", f"=SUMPRODUCT(ABS({lsr('xl_score_sumpts')}-{lsr('xl_score')}))",
               R("ref_zero"), "tol_pts"),
              ("Factor from PDO vs typed factor", f"={NM['factor_chk']}", f"={NM['factor']}", "tol_pts"),
              ("Offset from base score and odds vs typed offset", f"={NM['offset_chk']}", f"={NM['offset']}", "tol_pts"),
              ("Calculator: sum of points vs total score", "=" + addr["calc_points_sum"], "=" + addr["calc_score"], "tol_pts")]
    for k, (lab, live_f, py_f, tol) in enumerate(checks):
        r_ = 5 + k
        wx.cell(r_, 1, lab)
        wx.cell(r_, 2, live_f)
        wx.cell(r_, 3, py_f)
        wx.cell(r_, 4, f"=ABS(B{r_}-C{r_})")
        wx.cell(r_, 5, f"={NM[tol]}")
        wx.cell(r_, 6, f"=--(D{r_}<=E{r_})")
    c_last = 5 + len(checks) - 1
    wx["B2"] = f"=--(SUM(F5:F{c_last})=COUNT(F5:F{c_last}))"
    reg("checks_all", "Checks", "B2")
    wx.column_dimensions["A"].width = 58

    # ============================================================ Conclusions (static text from computed numbers)
    wn = ws["Conclusions"]
    p = lambda t: float(pop[t])  # noqa: E731
    lines = [
        "Conclusions (static text written by build_workbook.py from the computed numbers; not formulas)",
        "",
        f"Data: {config.DATA_SOURCE}. The workbook scores {n:,} loans issued in 2013, a grade-stratified sample of the out-of-time population.",
        f"Scorecard {rx.MODEL_NAME}: {len(variables)} WoE variables ({', '.join(variables)}), PDO {config.PDO}, base score {config.BASE_SCORE} at {config.BASE_ODDS}:1 odds.",
        "",
        f"Discrimination: on the full 2013 population the Gini is {p('gini'):.3f} and the KS is {p('ks'):.3f}, both in the Amber band "
        f"(Green needs Gini >= {config.THRESHOLDS['gini'][0]:g} and KS >= {config.THRESHOLDS['ks'][0]:g}). The sample deciles give a decile KS of {ref['ks_decile']:.3f} "
        f"and an approximate Gini of {ref['gini_approx']:.3f}; the exact sample Gini is {ref['gini']:.3f}. The decile Gini is an approximation, not a replacement for the exact figure.",
        "",
        f"Calibration: the sample mean PD is {cen['mean_pd']:.4f} against an observed 12-month bad rate of {cen['obs_rate']:.4f}, a ratio of {cen['ratio']:.2f}, "
        f"so the scorecard overstates risk in 2013. The sample Hosmer-Lemeshow statistic is {hl['stat']:.1f} (df {hl['df']}, p = {hl['p_value']:.3g}). "
        f"On the full population the HL p-value is {p('hl_p'):.3g} and the calibration slope is {p('cal_slope'):.3f}; the HL test has very high power at that size, "
        "so read it with the decile table.",
        "",
        f"Stability: the score PSI of the sample against the development population is {ref['psi_total']:.4f} (full population {p('psi_score'):.4f}), "
        f"below the {config.THRESHOLDS['psi'][0]:g} Green limit.",
        "",
        f"Expected loss: the live 12-month sample EL (PD x LGD x EAD) is ${ref['el'].sum():,.0f}; the Checks sheet reconciles every live figure to Python.",
        "",
        "Limits: this is an independent validation exercise on public data, not a bank-approved model. The workbook reproduces the scorecard, "
        "decile, HL, grade binomial, PSI and EL arithmetic; LGD and EAD models, the GBM benchmarks and bootstrap intervals are Python-only.",
    ]
    for k, t in enumerate(lines):
        wn.cell(1 + k, 1, t)
    wn["A1"].font = BOLD
    wn.column_dimensions["A"].width = 160

    # ============================================================ README
    wm = ws["README"]
    readme = [
        "Credit risk model validation workbook (LendingClub public 2007-2014 data)",
        "",
        "Purpose: a live, auditable Excel reproduction of the SC_FULL scorecard and its validation statistics on a grade-stratified sample of 2013 (OOT1) loans.",
        "Every number on Score_Calc, Deciles, Calibration, Checks and the live columns of Loan_Sample, PSI and EL is a formula. Typed numbers are limited to Inputs, Scorecard, Python_Ref and the typed Python columns.",
        "",
        "Sheets: Inputs (scorecard scaling, thresholds, tolerances, calculator attributes); Scorecard (points table); Score_Calc (one-loan calculator);",
        "Loan_Sample (attributes, outcomes, Python PD and points, live bin, WoE, points, score, PD, EL, ranks); Deciles (counts, bad rate, capture, lift, KS, approximate Gini);",
        "Calibration (Hosmer-Lemeshow, central tendency, Brier, grade z and Jeffreys tests); PSI (score stability); EL (expected vs realized loss by grade);",
        "Checks (live vs Python, cell B2 = 1 when all pass); Conclusions (static text); Python_Ref (Python results).",
        "",
        "Conventions: bins are (lo, hi] with a separate missing bin last (blank cells go there); unseen categories go to the largest group; WoE = ln(%good/%bad), so coefficients are negative;",
        "points = -(coef x WoE + intercept/k) x factor + offset/k; PD = 1/(1+EXP(-logit)); deciles rank descending by PD with ties broken by input order, decile = INT((rank-1)*10/N)+1.",
        "Bin lookups count edges below the value with SUMPRODUCT(--(edges<x))+1, a numeric comparison. COUNTIF(edges,\"<\"&x) gives the same bin except where text conversion at 15 digits "
        "moves a value that sits exactly on an edge, so the numeric form is used.",
        "",
        "Approximations: the decile Gini and KS are decile-based approximations (the exact loan-level figures are on Python_Ref). The calibration slope is a Python result, not a formula.",
        "The sample statistics differ from the full-population figures in validation_results.csv (Python_Ref block B) because the sample is stratified and smaller.",
        "Author: Deepak Chaudhary. Not a bank-approved model.",
    ]
    for k, t in enumerate(readme):
        wm.cell(1 + k, 1, t)
    wm["A1"].font = BOLD
    wm.column_dimensions["A"].width = 160

    wb.save(out_file)
    n_formulas = sum(1 for s in wb for rw in s.iter_rows() for c in rw if isinstance(c.value, str) and c.value.startswith("="))
    return dict(addr=addr, first=FIRST, last=last, hdr=HDR, ls_cols=lc, n_formulas=n_formulas, n_loans=n,
                deciles=dict(first=D0), hl=dict(first=K0), grades=dict(first=G0, labels=grades), psi=dict(first=P0, last=P1),
                el=dict(first=E0, labels=grades), calc=dict(first=CF), sheets=names,
                typed_cols=dict(PSI=list("ABCDEFG"), EL=list("AHIJ")), variables=variables)
